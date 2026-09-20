from __future__ import annotations

from io import BytesIO
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SUCCESS_PAYMENT_STATUSES = {
    "paid",
    "succeeded",
}

SUCCESS_REFUND_STATUSES = {
    "succeeded",
}

CLOSED_BOOKING_STATUSES = {
    "Отмена",
    "Не актуальна",
}


def _money(value) -> float:
    return float(
        Decimal(value or 0)
    )


def _excel_datetime(value):
    """Excel/openpyxl does not accept timezone-aware datetimes."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone().replace(tzinfo=None)
    return value


def _payment_refunded_total(
    payment,
) -> Decimal:
    return sum(
        (
            Decimal(refund.amount or 0)
            for refund in payment.refunds
            if refund.status
            in SUCCESS_REFUND_STATUSES
        ),
        Decimal("0"),
    )


def _booking_gross_paid(
    booking,
) -> Decimal:
    return sum(
        (
            Decimal(payment.amount or 0)
            for payment in booking.payments
            if payment.status
            in SUCCESS_PAYMENT_STATUSES
        ),
        Decimal("0"),
    )


def _booking_refunded(
    booking,
) -> Decimal:
    return sum(
        (
            _payment_refunded_total(payment)
            for payment in booking.payments
            if payment.status
            in SUCCESS_PAYMENT_STATUSES
        ),
        Decimal("0"),
    )


def _booking_net_paid(
    booking,
) -> Decimal:
    return max(
        _booking_gross_paid(booking)
        - _booking_refunded(booking),
        Decimal("0"),
    )


def _applicant_name(booking) -> str:
    return (
        booking.applicant_full_name
        or (
            booking.customer.full_name
            if booking.customer
            else ""
        )
        or ""
    )


def _phone(booking) -> str:
    return (
        booking.applicant_phone
        or (
            booking.customer.phone
            if booking.customer
            else ""
        )
        or ""
    )


def _telegram(booking) -> str:
    username = (
        booking.applicant_telegram_username
        or (
            booking.customer.telegram_username
            if booking.customer
            else ""
        )
        or ""
    )

    username = str(username).strip()

    if username and not username.startswith("@"):
        username = f"@{username}"

    return username


def _participant_accommodation(
    participant,
) -> str:
    if participant.accommodation_rate:
        return (
            participant.accommodation_rate.title
            or ""
        )

    return (
        participant.accommodation_type
        or "Не выбрано"
    )


def _participant_room(
    participant,
) -> str:
    if participant.room:
        return str(
            participant.room.number
            or ""
        )

    return "Не расселён"


def _format_sheet(
    sheet,
    *,
    freeze: str = "A2",
):
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions

    header_fill = PatternFill(
        "solid",
        fgColor="1F2937",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    for row in sheet.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    widths = {}

    for row in sheet.iter_rows():
        for cell in row:
            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            widths[cell.column] = min(
                max(
                    widths.get(
                        cell.column,
                        0,
                    ),
                    len(value) + 2,
                ),
                38,
            )

    for column_index, width in (
        widths.items()
    ):
        sheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = max(
            width,
            10,
        )


def build_departure_xlsx(
    departure,
) -> BytesIO:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Сводка"

    active_bookings = [
        booking
        for booking in departure.bookings
        if booking.status
        not in CLOSED_BOOKING_STATUSES
    ]

    active_participants = [
        participant
        for booking in active_bookings
        for participant
        in booking.participants
    ]

    tickets_bought = sum(
        1
        for participant
        in active_participants
        if participant.ticket_status
        == "Куплен"
    )

    tickets_missing = (
        len(active_participants)
        - tickets_bought
    )

    gross_paid = sum(
        (
            _booking_gross_paid(
                booking
            )
            for booking
            in departure.bookings
        ),
        Decimal("0"),
    )

    refunded = sum(
        (
            _booking_refunded(
                booking
            )
            for booking
            in departure.bookings
        ),
        Decimal("0"),
    )

    net_paid = max(
        gross_paid - refunded,
        Decimal("0"),
    )

    summary_rows = [
        (
            "Параметр",
            "Значение",
        ),
        (
            "Направление",
            departure.tour.title,
        ),
        (
            "Начало",
            departure.start_date,
        ),
        (
            "Окончание",
            departure.end_date,
        ),
        (
            "Сезон",
            getattr(
                departure,
                "season",
                "",
            )
            or "",
        ),
        (
            "Вместимость",
            departure.capacity,
        ),
        (
            "Активных заявок",
            len(active_bookings),
        ),
        (
            "Участников в активных заявках",
            len(active_participants),
        ),
        (
            "Билетов куплено",
            tickets_bought,
        ),
        (
            "Билетов нужно купить",
            tickets_missing,
        ),
        (
            "Поступило, ₽",
            _money(gross_paid),
        ),
        (
            "Возвращено, ₽",
            _money(refunded),
        ),
        (
            "Чистыми, ₽",
            _money(net_paid),
        ),
    ]

    for row in summary_rows:
        summary.append(row)

    _format_sheet(
        summary,
        freeze="A2",
    )

    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 28

    for cell in (
        summary["B11"],
        summary["B12"],
        summary["B13"],
    ):
        cell.number_format = (
            '#,##0.00 "₽"'
        )

    participants = workbook.create_sheet(
        "Участники"
    )

    participants.append(
        [
            "Заявка",
            "Статус заявки",
            "ФИО участника",
            "ФИО заявителя",
            "Телефон",
            "Telegram",
            "Источник",
            "Размещение",
            "Комната",
            "Билет",
            "Стоимость участника, ₽",
            "Комментарий участника",
            "Комментарий заявки",
        ]
    )

    bookings_sorted = sorted(
        departure.bookings,
        key=lambda booking: (
            booking.created_at,
            booking.id,
        ),
    )

    for booking in bookings_sorted:
        public_number = (
            f"WT-{booking.id:04d}"
        )

        for participant in sorted(
            booking.participants,
            key=lambda item: item.id,
        ):
            participants.append(
                [
                    public_number,
                    booking.status,
                    participant.full_name,
                    _applicant_name(
                        booking
                    ),
                    _phone(booking),
                    _telegram(booking),
                    booking.source or "",
                    _participant_accommodation(
                        participant
                    ),
                    _participant_room(
                        participant
                    ),
                    (
                        participant.ticket_status
                        or "Не куплен"
                    ),
                    (
                        _money(
                            participant
                            .accommodation_price
                        )
                        if participant
                        .accommodation_price
                        is not None
                        else ""
                    ),
                    participant.comment or "",
                    booking.comment or "",
                ]
            )

    _format_sheet(
        participants
    )

    for row in participants.iter_rows(
        min_row=2,
        min_col=11,
        max_col=11,
    ):
        row[0].number_format = (
            '#,##0.00 "₽"'
        )

    bookings_sheet = (
        workbook.create_sheet(
            "Заявки"
        )
    )

    bookings_sheet.append(
        [
            "Заявка",
            "Дата создания",
            "Статус",
            "ФИО заявителя",
            "Телефон",
            "Telegram",
            "Источник",
            "Людей",
            "Стоимость, ₽",
            "Поступило, ₽",
            "Возвращено, ₽",
            "Чистыми, ₽",
            "Осталось, ₽",
            "Комментарий",
        ]
    )

    for booking in bookings_sorted:
        gross = _booking_gross_paid(
            booking
        )
        booking_refunded = (
            _booking_refunded(
                booking
            )
        )
        net = max(
            gross - booking_refunded,
            Decimal("0"),
        )
        agreed = Decimal(
            booking.agreed_price
            or 0
        )
        remaining = max(
            agreed - net,
            Decimal("0"),
        )

        bookings_sheet.append(
            [
                f"WT-{booking.id:04d}",
                _excel_datetime(booking.created_at),
                booking.status,
                _applicant_name(
                    booking
                ),
                _phone(booking),
                _telegram(booking),
                booking.source or "",
                booking.people_count,
                _money(agreed),
                _money(gross),
                _money(
                    booking_refunded
                ),
                _money(net),
                _money(remaining),
                booking.comment or "",
            ]
        )

    _format_sheet(
        bookings_sheet
    )

    for column in (
        9,
        10,
        11,
        12,
        13,
    ):
        for row in bookings_sheet.iter_rows(
            min_row=2,
            min_col=column,
            max_col=column,
        ):
            row[0].number_format = (
                '#,##0.00 "₽"'
            )

    payments_sheet = (
        workbook.create_sheet(
            "Платежи"
        )
    )

    payments_sheet.append(
        [
            "Заявка",
            "Дата",
            "Тип записи",
            "Сумма, ₽",
            "Способ",
            "Статус",
            "Комментарий",
        ]
    )

    for booking in bookings_sorted:
        public_number = (
            f"WT-{booking.id:04d}"
        )

        for payment in sorted(
            booking.payments,
            key=lambda item: (
                item.created_at,
                item.id,
            ),
        ):
            payment_date = (
                payment.paid_at
                or payment.created_at
            )

            payments_sheet.append(
                [
                    public_number,
                    _excel_datetime(payment_date),
                    "Платёж",
                    _money(
                        payment.amount
                    ),
                    payment.provider or "",
                    payment.status or "",
                    payment.comment or "",
                ]
            )

            for refund in sorted(
                payment.refunds,
                key=lambda item: (
                    item.created_at,
                    item.id,
                ),
            ):
                refund_date = (
                    refund.refunded_at
                    or refund.created_at
                )

                payments_sheet.append(
                    [
                        public_number,
                        _excel_datetime(refund_date),
                        "Возврат",
                        -_money(
                            refund.amount
                        ),
                        (
                            refund.refund_method
                            or ""
                        ),
                        refund.status or "",
                        refund.comment or "",
                    ]
                )

    _format_sheet(
        payments_sheet
    )

    for row in payments_sheet.iter_rows(
        min_row=2,
        min_col=4,
        max_col=4,
    ):
        row[0].number_format = (
            '#,##0.00 "₽"'
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output
